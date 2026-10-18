"""Public synthetic tests: brand onboarding, independent export, and shared defaults."""
import base64
import io
import json
import sys
import tempfile
import unittest
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
GENERIC = ROOT / 'skills/wechat-brand-layout/scripts'
sys.path.insert(0, str(GENERIC))
import brand_layout as brand



class LayoutWorkflows(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.article = {'version': 1, 'title': '活动回顾', 'column': '文化活动',
                        'subtype': '活动总结', 'classification_reviewed': True,
                        'source_text': '活动回顾\n现场交流\n本次活动围绕实践与分享展开。',
                        'summary': '本次活动的实践与分享。', 'blocks': [
                            {'type': 'heading', 'level': 2, 'runs': [{'text': '现场交流'}]},
                            {'type': 'group', 'variant': 'section', 'blocks': [
                                {'type': 'paragraph', 'runs': [{'text': '本次活动围绕实践与分享展开。'}]}]}]}
        self.input = self.root / 'article.json'
        brand.write_json(self.input, self.article)

    def tearDown(self):
        self.temp.cleanup()

    def profile(self, name='example'):
        return Path(brand.init_profile(self.root / 'data', name, '测试品牌 ' + name)['profile'])

    def image(self, name, color):
        path = self.root / name
        Image.new('RGB', (320, 100), color).save(path, format='PNG')
        return path

    def test_first_use_requires_response_and_profiles_do_not_overwrite(self):
        profile = self.profile()
        with self.assertRaisesRegex(ValueError, '首次素材征集'):
            brand.render_article(self.input, profile, self.root / 'out')
        with self.assertRaisesRegex(ValueError, '已存在'):
            self.profile()
        brand.mark_ready(profile, 'neutral')
        result = brand.render_article(self.input, profile, self.root / 'out')
        self.assertTrue(result['xiumi_package_generated'])
        packet = brand.read_json(self.root / 'out/article.xiumi.json')
        self.assertEqual(packet['assetSummary']['embedded'], 0)
        self.assertEqual(packet['schema'], 'xiumi-bridge/v3')
        self.assertNotIn('吴', packet['html'])
        self.assertEqual(brand.list_profiles(self.root / 'data')[0]['status'], 'ready')

    def test_assets_persist_roles_are_respected_and_rerender_is_idempotent(self):
        profile = self.profile()
        for name, color, role in [('top.bin', 'red', 'header'), ('bottom.png', 'blue', 'footer'),
                                   ('cover.png', 'green', 'cover'), ('logo.png', 'yellow', 'logo')]:
            source = self.image(name, color)
            asset = brand.import_assets(profile, source)['imported'][0]
            brand.assign_asset(profile, asset['id'], role, role in ('header', 'footer'))
            source.unlink()  # Chat attachments can disappear; permanent copied assets must work.
        brand.mark_ready(profile, 'configured')
        out = self.root / 'out'
        brand.render_article(self.input, profile, out)
        packet = brand.read_json(out / 'article.xiumi.json')
        self.assertEqual(packet['assetSummary']['embedded'], 2)
        self.assertEqual(packet['assetSummary']['local'], 0)
        for asset in packet['assets']:
            prefix, data = asset['src'].split(',', 1)
            self.assertEqual(prefix, 'data:image/png;base64')
            with Image.open(io.BytesIO(base64.b64decode(data))) as img:
                self.assertIn(img.getpixel((0, 0)), [(255, 0, 0), (0, 0, 255)])
        brand.render_article(out / 'article.json', profile, self.root / 'again')
        self.assertEqual(len(brand.read_json(self.root / 'again/article.xiumi.json')['native']['modules']),
                         len(packet['native']['modules']))

    def test_brand_isolation_and_style_override(self):
        p1, p2 = self.profile('one'), self.profile('two')
        for path, color in [(p1, '#287746'), (p2, '#ac3f44')]:
            profile = brand.read_json(path)
            profile['theme']['primary'] = color
            profile['theme']['group_border'] = None
            brand.write_json(path, profile)
            brand.mark_ready(path, 'neutral')
        brand.render_article(self.input, p1, self.root / 'one')
        brand.render_article(self.input, p2, self.root / 'two')
        for name, color in [('one', '#287746'), ('two', '#ac3f44')]:
            html = (self.root / name / 'article.html').read_text(encoding='utf-8')
            self.assertIn('color:' + color, html)
            self.assertIn('padding:22px', html)
        with self.assertRaisesRegex(ValueError, '其他账号'):
            brand.render_article(self.root / 'one/article.json', p2, self.root / 'bad')

    def test_zip_dedup_unsafe_names_and_unsupported_images(self):
        profile = self.profile()
        logo = self.image('logo.png', 'orange').read_bytes()
        archive = self.root / 'assets.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('../outside.png', logo)
            z.writestr('second/logo.png', logo)
            z.writestr('drawing.svg', '<svg/>')
        result = brand.import_assets(profile, archive)
        self.assertTrue(result['imported'][1]['duplicate'])
        self.assertFalse((self.root / 'outside.png').exists())
        self.assertEqual(len(brand.read_json(profile)['assets']), 2)
        with self.assertRaisesRegex(ValueError, '不能直接渲染'):
            brand.assign_asset(profile, result['imported'][2]['id'], 'footer', True)
        with self.assertRaisesRegex(ValueError, '仅正文开头'):
            brand.assign_asset(profile, result['imported'][0]['id'], 'cover', True)

    def test_missing_image_or_changed_text_cannot_export(self):
        profile = self.profile()
        brand.mark_ready(profile, 'neutral')
        a = deepcopy(self.article)
        a['blocks'][1]['blocks'][0]['runs'][0]['text'] = '删改了正文。'
        brand.write_json(self.input, a)
        with self.assertRaisesRegex(ValueError, '内容与提取稿不同'):
            brand.render_article(self.input, profile, self.root / 'changed')
        a = deepcopy(self.article)
        a['blocks'].append({'type': 'image', 'src': 'missing.png'})
        brand.write_json(self.input, a)
        with self.assertRaisesRegex(ValueError, '图片不存在'):
            brand.render_article(self.input, profile, self.root / 'missing')
        self.assertFalse((self.root / 'missing/article.xiumi.json').exists())

    def test_adjacent_gallery_keeps_images_independent(self):
        from bs4 import BeautifulSoup
        article = deepcopy(self.article)
        article['blocks'].append({'type': 'gallery', 'mode': 'row', 'images': [
            {'src': str(self.image(f'photo-{i}.png', color))}
            for i, color in enumerate(['red', 'blue', 'green'])]})
        profile = self.profile()
        brand.mark_ready(profile, 'neutral')
        brand.write_json(self.input, article)
        brand.render_article(self.input, profile, self.root / 'generic-row')
        for folder in ['generic-row']:
            packet = brand.read_json(self.root / folder / 'article.xiumi.json')
            row = BeautifulSoup(packet['html'], 'html.parser').select_one('[data-image-layout="row"]')
            self.assertEqual(len(row.select('td:has(img)')), 3)
            self.assertEqual(len(row.select('[data-image-gap]')), 2)
            self.assertTrue(all('width:8px' in gap['style'] for gap in row.select('[data-image-gap]')))
            self.assertEqual(len(set(img['src'] for img in row.select('img'))), 3)
            self.assertEqual(packet['assetSummary']['embedded'], 3)
            self.assertEqual(packet['assetSummary']['local'], 0)

    def test_docx_with_and_without_numbering_part(self):
        original = self.root / 'normal.docx'
        doc = Document()
        doc.add_heading('测试通讯稿', 0)
        doc.add_paragraph('原文包含重点与一张图片。')
        doc.add_picture(str(self.image('photo.png', 'cyan')))
        doc.save(original)
        no_numbering = self.root / 'no-numbering.docx'
        with zipfile.ZipFile(original) as src, zipfile.ZipFile(no_numbering, 'w') as dst:
            for item in src.infolist():
                if item.filename == 'word/numbering.xml':
                    continue
                data = src.read(item)
                if item.filename in ('word/_rels/document.xml.rels', '[Content_Types].xml'):
                    xml = etree.fromstring(data)
                    for node in list(xml):
                        if 'numbering' in node.get('Type', '') or node.get('PartName') == '/word/numbering.xml':
                            xml.remove(node)
                    data = etree.tostring(xml)
                dst.writestr(item, data)
        for index, module in enumerate((brand.engine,)):
            a = module.extract(no_numbering, self.root / f'work-{index}/article.json')
            self.assertIn('原文包含重点', a['source_text'])
            self.assertEqual(sum(b['type'] == 'image' for b in a['blocks']), 1)
        doc.add_paragraph('第一项', style='List Number')
        doc.save(original)
        for index, module in enumerate((brand.engine,)):
            a = module.extract(original, self.root / f'numbered-{index}/article.json')
            self.assertTrue(any(b['type'] == 'list_item' for b in a['blocks']))


if __name__ == '__main__':
    unittest.main()
